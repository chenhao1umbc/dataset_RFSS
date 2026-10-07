"""Build figures/fig_pipeline.pptx, the source of Figure 1 (dataset construction pipeline).

The slide is drawn at the final print size (7.16 in wide, the IEEEtran two-column text width), so the exported PDF needs no
rescaling in LaTeX. Export to figures/fig_pipeline.pdf from Keynote (File > Export To > PDF) or PowerPoint, as for the AAMAS2027
figure; the blocks stay editable in the .pptx.

    uv run --no-project --with python-pptx python gen_fig_pipeline.py
"""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent / "figures" / "fig_pipeline.pptx"
WIDTH_IN, HEIGHT_IN = 7.16, 2.50
FONT = "Times New Roman"
SIZE_BODY, SIZE_SMALL = 7.0, 7.0
GREEN_BAR = RGBColor(0x1F, 0x6E, 0x30)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY_TEXT = RGBColor(0x55, 0x55, 0x55)
DASH_GREEN = RGBColor(0x2E, 0x8B, 0x3E)
FILL = {
    "symbols": (RGBColor(0xEE, 0xEE, 0xEE), RGBColor(0x9A, 0x9A, 0x9A)),
    "generator": (RGBColor(0xE6, 0xEC, 0xF8), RGBColor(0x8C, 0x9C, 0xC4)),
    "channel": (RGBColor(0xE2, 0xF1, 0xF5), RGBColor(0x7F, 0xB4, 0xC2)),
    "impair": (RGBColor(0xFD, 0xEF, 0xDC), RGBColor(0xD9, 0xA4, 0x66)),
    "mixer": (RGBColor(0xE6, 0xF3, 0xE3), RGBColor(0x6E, 0xAE, 0x6A)),
    "noise": (RGBColor(0xEE, 0xEE, 0xEE), RGBColor(0x9A, 0x9A, 0x9A)),
    "mixture": (RGBColor(0xDC, 0xE9, 0xFB), RGBColor(0x2F, 0x6F, 0xD6)),
    "refs": (RGBColor(0xE6, 0xF3, 0xE3), RGBColor(0x6E, 0xAE, 0x6A)),
    "meta": (RGBColor(0xEE, 0xEE, 0xEE), RGBColor(0x9A, 0x9A, 0x9A)),
}


def emu(value):
    return Emu(int(Inches(value)))


def add_text(frame, lines, align=PP_ALIGN.CENTER):
    """lines: (text, bold, italic, grey) tuples, one paragraph each, all at the body size."""
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = emu(0.03)
    frame.margin_top = frame.margin_bottom = emu(0.02)
    for index, (text, bold, italic, grey) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = 0.95
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(SIZE_BODY)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = GREY_TEXT if grey else INK


def add_block(slide, x, y, w, h, kind, lines, line_width=0.75):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h)
    )
    shape.adjustments[0] = 0.08
    fill, edge = FILL[kind]
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = edge
    shape.line.width = Pt(line_width)
    shape.shadow.inherit = False
    add_text(shape.text_frame, lines)
    return shape


def add_bar(slide, x, y, w, h, text):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h)
    )
    shape.adjustments[0] = 0.35
    shape.fill.solid()
    shape.fill.fore_color.rgb = GREEN_BAR
    shape.line.fill.background()
    shape.shadow.inherit = False
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_top = frame.margin_bottom = emu(0.0)
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(8.0)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    frame.paragraphs[0].alignment = PP_ALIGN.CENTER


def add_label(slide, x, y, w, h, text, italic=True, align=PP_ALIGN.CENTER):
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    add_text(box.text_frame, [(text, False, italic, True)], align)


def add_line(slide, x1, y1, x2, y2, dashed=False, arrow=True):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2)
    )
    connector.line.width = Pt(1.0 if dashed else 0.75)
    connector.line.color.rgb = DASH_GREEN if dashed else INK
    if dashed:
        connector.line.dash_style = MSO_LINE.DASH
    if arrow:
        outline = connector.line._get_or_add_ln()
        etree.SubElement(outline, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return connector


def build():
    presentation = Presentation()
    presentation.slide_width = emu(WIDTH_IN)
    presentation.slide_height = emu(HEIGHT_IN)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    add_bar(
        slide,
        0.05,
        0.04,
        4.30,
        0.22,
        "(a) Per-source chain, applied to each of the K sources",
    )
    add_bar(slide, 4.45, 0.04, 1.40, 0.22, "(b) Mixing")
    add_bar(slide, 5.97, 0.04, 1.14, 0.22, "(c) HDF5 output")
    add_label(
        slide,
        0.05,
        0.28,
        4.30,
        0.30,
        "K = 2, 3 or 4 sources per sample, drawn with replacement; each source has its own chain",
    )

    top, height = 0.72, 1.00
    chain = [
        (
            "symbols",
            0.12,
            0.58,
            [("Random", True, False, False), ("symbols", True, False, False)],
        ),
        (
            "generator",
            0.86,
            0.98,
            [
                ("Generator", True, False, False),
                ("GSM, UMTS,", False, False, False),
                ("LTE, 5G NR", False, False, False),
                ("TS 45.004, 25.213,", False, False, True),
                ("36.211, 38.211", False, False, True),
            ],
        ),
        (
            "channel",
            2.00,
            0.82,
            [
                ("TDL channel", True, False, False),
                ("TDL-A to TDL-E", False, False, False),
                ("TR 38.901", False, False, True),
            ],
        ),
        (
            "impair",
            2.98,
            1.30,
            [
                ("Hardware impairments", True, False, False),
                ("CFO, SFO, I/Q imbalance,", False, False, False),
                ("DC offset, phase noise,", False, False, False),
                ("PA nonlinearity", False, False, False),
                ("clean 19.9%, single 29.9%,", False, False, True),
                ("all six 50.1%", False, False, True),
            ],
        ),
    ]
    for kind, x, w, lines in chain:
        for offset in (0.08, 0.04):
            shadow = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                emu(x + offset),
                emu(top + offset),
                emu(w),
                emu(height),
            )
            shadow.adjustments[0] = 0.08
            shadow.fill.solid()
            shadow.fill.fore_color.rgb = RGBColor(0xFA, 0xFA, 0xFA)
            shadow.line.color.rgb = FILL[kind][1]
            shadow.line.width = Pt(0.5)
            shadow.shadow.inherit = False
        add_block(slide, x, top, w, height, kind, lines)
    for x_from, x_to in ((0.70, 0.86), (1.84, 2.00), (2.82, 2.98)):
        add_line(slide, x_from, top + height / 2, x_to, top + height / 2)

    mixer_x, mixer_w = 4.55, 1.22
    add_block(
        slide,
        mixer_x,
        top,
        mixer_w,
        height,
        "mixer",
        [
            ("SignalMixer", True, False, False),
            ("rate = highest", False, False, False),
            ("source rate", False, False, False),
            ("power ratio per source", False, False, False),
            ("offset 0 (co-channel)", False, False, False),
            ("or k x 2 MHz (adjacent)", False, False, False),
        ],
    )
    add_line(slide, 4.38, top + height / 2, mixer_x, top + height / 2)
    awgn_top = 1.84
    add_block(
        slide,
        mixer_x,
        awgn_top,
        mixer_w,
        0.36,
        "noise",
        [
            ("+ AWGN, sample SNR", True, False, False),
            ("against mixture power", False, False, True),
        ],
    )
    add_line(
        slide, mixer_x + mixer_w / 2, top + height, mixer_x + mixer_w / 2, awgn_top
    )

    out_x, out_w, out_h = 6.10, 1.00, 0.44
    rows = {"mixture": 0.60, "refs": 1.14, "meta": 1.68}
    add_block(
        slide,
        out_x,
        rows["mixture"],
        out_w,
        out_h,
        "mixture",
        [("Mixture", True, False, False), ("one array per sample", False, False, True)],
        1.25,
    )
    add_block(
        slide,
        out_x,
        rows["refs"],
        out_w,
        out_h,
        "refs",
        [
            ("Per-source references", True, False, False),
            ("stored before mixing", False, False, True),
        ],
    )
    add_block(
        slide,
        out_x,
        rows["meta"],
        out_w,
        out_h,
        "meta",
        [("Metadata", True, False, False), ("JSON per sample", False, False, True)],
    )

    solid_x = 5.89
    awgn_mid = awgn_top + 0.18
    mixture_mid = rows["mixture"] + out_h / 2
    add_line(slide, mixer_x + mixer_w, awgn_mid, solid_x, awgn_mid, arrow=False)
    add_line(slide, solid_x, awgn_mid, solid_x, mixture_mid, arrow=False)
    add_line(slide, solid_x, mixture_mid, out_x, mixture_mid)

    corridor_y, dashed_x = 2.30, 6.00
    tap_x = 3.63
    refs_mid, meta_mid = rows["refs"] + out_h / 2, rows["meta"] + out_h / 2
    add_line(slide, tap_x, top + height, tap_x, corridor_y, dashed=True, arrow=False)
    add_line(slide, tap_x, corridor_y, dashed_x, corridor_y, dashed=True, arrow=False)
    add_line(slide, dashed_x, corridor_y, dashed_x, refs_mid, dashed=True, arrow=False)
    add_line(slide, dashed_x, refs_mid, out_x, refs_mid, dashed=True)
    add_line(slide, dashed_x, meta_mid, out_x, meta_mid, dashed=True)
    add_label(slide, 4.15, 2.31, 1.70, 0.16, "taps leave the chain before the mixer")

    presentation.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
